"""Step 2: fetch every candidate twice (plain HTTP and a real browser render) and cache what later steps need.

  02_fetch.py [--limit N] [--concurrency 8] [--urls-from FILE]

Cache per URL (out/cache/xx/<sha1>.*): .json.gz (blocks, raw text, signals), .jpg (viewport screenshot), .html.gz.
"""
import argparse
import asyncio
import datetime as dt
import gzip
import io
import random
import re
import sys
import urllib.robotparser
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from selectolax.lexbor import LexborHTMLParser as HTMLParser

from common import cache_path, db, write_gz_json, smart_tokenize

RAW_UA = "Mozilla/5.0 (compatible; huntsman-scrape-dataset-v1/1.0)"
BROWSER_UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/{v} Safari/537.36"
ROBOTS_UA = "huntsman-scrape-dataset-v1"
EXTRACT_JS = (Path(__file__).with_name("extract.js")).read_text()
MAX_BYTES = 25 * 1024 * 1024
CHALLENGE_RE = re.compile(r"just a moment|verify you are (a )?human|attention required|checking your browser|"
                          r"access denied|are you a robot|captcha|ddos protection|please enable (javascript and )?cookies|"
                          r"request unsuccessful\. incapsula|pardon our interruption|bot detection", re.I)
BLOCK_MARKERS = ("cf-chl", "challenge-platform", "_incapsula_resource", "datadome", "perimeterx", "px-captcha",
                 "akamai-bot", "sucuri", "ak_bmsc")
OFFICE_EXT = {".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx"}
LOCALE_RE = re.compile(r"(^|[/.?=&_-])([a-z]{2}([-_][a-z]{2})?)([/.&]|$)", re.I)


def doc_kind(url, content_type):
    ct = (content_type or "").lower()
    path = urlsplit(url).path.lower()
    if "pdf" in ct or path.endswith(".pdf"):
        return "pdf"
    if "wordprocessingml" in ct or path.endswith(".docx"):
        return "docx"
    if "presentationml" in ct or path.endswith(".pptx"):
        return "pptx"
    if "spreadsheetml" in ct or path.endswith(".xlsx"):
        return "xlsx"
    return None


def doc_blocks(kind, data):
    """Text blocks for PDF/Office files, plus whether the PDF looks scanned (no text layer)."""
    blocks, scanned = [], False
    if kind == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = reader.pages[:60]
        texts = [(p.extract_text() or "") for p in pages]
        scanned = sum(len(t.strip()) for t in texts) < 100 * max(len(pages), 1) * 0.3
        for n, t in enumerate(texts, 1):
            for para in re.split(r"\n\s*\n|(?<=[.!?。])\n", t):
                para = re.sub(r"\s+", " ", para).strip()
                if para:
                    blocks.append({"tag": "p", "ctx": f"page{n}", "text": para[:5000]})
    elif kind == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        for sec in d.sections:
            for p in sec.header.paragraphs:
                if p.text.strip():
                    blocks.append({"tag": "p", "ctx": "header", "text": p.text.strip()})
        for p in d.paragraphs:
            if p.text.strip():
                blocks.append({"tag": "h" if p.style.name.lower().startswith("heading") else "p", "ctx": "body", "text": p.text.strip()})
        for t in d.tables:
            for row in t.rows:
                txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if txt:
                    blocks.append({"tag": "tr", "ctx": "body>table", "text": txt})
        for sec in d.sections:
            for p in sec.footer.paragraphs:
                if p.text.strip():
                    blocks.append({"tag": "p", "ctx": "footer", "text": p.text.strip()})
    elif kind == "pptx":
        import pptx
        pres = pptx.Presentation(io.BytesIO(data))
        for n, slide in enumerate(pres.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = "".join(r.text for r in p.runs).strip()
                        if txt:
                            blocks.append({"tag": "p", "ctx": f"slide{n}", "text": txt})
    elif kind == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets[:5]:
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                if r >= 400:
                    break
                txt = " | ".join(str(v) for v in row if v not in (None, ""))
                if txt:
                    blocks.append({"tag": "tr", "ctx": f"sheet:{ws.title}", "text": txt[:2000]})
    for i, b in enumerate(blocks):
        b.update(i=i, shadow=False, iframe=False)
    return blocks, scanned


def raw_html_text(html):
    tree = HTMLParser(html)
    for node in tree.css("script, style, noscript, template, svg"):
        node.decompose()
    body = tree.body or tree.root
    return re.sub(r"\s+", " ", body.text(separator=" ") if body else "").strip()


def detect_charset(resp, data):
    ct = resp.headers.get("content-type", "")
    m = re.search(r"charset=([\w-]+)", ct, re.I)
    if m:
        return m.group(1).lower()
    m = re.search(rb"<meta[^>]+charset=[\"']?([\w-]+)", data[:4096], re.I)
    return m.group(1).decode().lower() if m else None


class Fetcher:
    def __init__(self, browser, client):
        self.browser, self.client = browser, client
        self.robots = {}

    async def allowed(self, url):
        parts = urlsplit(url)
        origin = f"{parts.scheme}://{parts.netloc}"
        if origin not in self.robots:
            rp = urllib.robotparser.RobotFileParser()
            try:
                r = await self.client.get(origin + "/robots.txt", timeout=10)
                rp.parse(r.text.splitlines() if r.status_code == 200 else [])
            except Exception:
                rp.parse([])
            self.robots[origin] = rp
        return self.robots[origin].can_fetch(ROBOTS_UA, url)

    async def raw(self, url):
        out = {"hops": [], "status": None, "content_type": None, "charset": None, "bytes": 0}
        try:
            async with self.client.stream("GET", url, timeout=25) as resp:
                chunks, size = [], 0
                async for c in resp.aiter_bytes():
                    chunks.append(c)
                    size += len(c)
                    if size > MAX_BYTES:
                        break
                data = b"".join(chunks)
                out.update(status=resp.status_code, final_url=str(resp.url), content_type=resp.headers.get("content-type"),
                           hops=[str(h.url) for h in resp.history], bytes=size, charset=detect_charset(resp, data),
                           server=resp.headers.get("server"))
                return out, data
        except Exception as e:
            out["error"] = f"{type(e).__name__}: {str(e)[:200]}"
            return out, b""

    async def render(self, url):
        ctx = await self.browser.new_context(user_agent=BROWSER_UA.format(v=self.browser.version), viewport={"width": 1366, "height": 900}, ignore_https_errors=True,
                                             java_script_enabled=True)
        page = await ctx.new_page()
        # text is all we need: skip heavy assets so pages settle sooner (screenshots show image placeholders)
        await page.route("**/*", lambda route: route.abort() if route.request.resource_type in ("image", "media", "font")
                         else route.continue_())
        nav_hops = []
        page.on("response", lambda r: nav_hops.append(r.status) if r.request.is_navigation_request() and r.frame == page.main_frame else None)
        try:
            resp = await page.goto(url, wait_until="domcontentloaded", timeout=25000)
            try:
                await page.wait_for_load_state("networkidle", timeout=3000)
            except Exception:
                pass
            before = await page.evaluate("document.body ? document.body.innerText.length : 0")
            for _ in range(12):
                at_bottom = await page.evaluate("(() => { window.scrollBy(0, window.innerHeight); return window.innerHeight + window.scrollY >= document.body.scrollHeight - 5; })()")
                await page.wait_for_timeout(200)
                if at_bottom:
                    break
            await page.wait_for_timeout(600)
            after = await page.evaluate("document.body ? document.body.innerText.length : 0")
            await page.evaluate("window.scrollTo(0, 0)")
            await page.wait_for_timeout(150)
            data = await page.evaluate(EXTRACT_JS)
            # cross-origin iframes are invisible to the in-page walk; read them through Playwright
            for frame in page.frames[1:8]:
                try:
                    sub = await frame.evaluate("document.body ? document.body.innerText : ''")
                except Exception:
                    continue
                for para in re.split(r"\n\s*\n", sub or ""):
                    para = re.sub(r"\s+", " ", para).strip()
                    if para and len(data["blocks"]) < 4000:
                        data["blocks"].append({"i": len(data["blocks"]), "tag": "text", "ctx": "iframe", "text": para[:5000],
                                               "shadow": False, "iframe": True})
                        data["iframeChars"] += len(para)
                        data["totalChars"] += len(para)
            shot = await page.screenshot(type="jpeg", quality=55)
            html = await page.content()
            data.update(final_url=page.url, status=resp.status if resp else None, text_before_scroll=before,
                        text_after_scroll=after, nav_statuses=nav_hops)
            return data, shot, html
        finally:
            await ctx.close()

    async def fetch(self, url):
        now = dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")
        if not await self.allowed(url):
            return {"ok": 0, "reason": "robots", "fetched_at": now, "meta": {}}
        raw, data = await self.raw(url)
        meta = {"raw": raw}
        kind = doc_kind(raw.get("final_url") or url, raw.get("content_type"))
        if kind:
            if not raw.get("status") or raw["status"] >= 400 or not data:
                return {"ok": 0, "reason": f"doc_http_{raw.get('status')}", "fetched_at": now, "meta": meta}
            try:
                blocks, scanned = await asyncio.to_thread(doc_blocks, kind, data)
            except Exception as e:
                return {"ok": 0, "reason": f"doc_parse: {str(e)[:100]}", "fetched_at": now, "meta": meta}
            text = "\n".join(b["text"] for b in blocks)
            meta.update(kind=kind, scanned=scanned, final_url=raw.get("final_url"), http_status=raw["status"])
            write_gz_json(cache_path(url, "json.gz"), {"kind": kind, "blocks": blocks, "raw_text": text, "meta": meta})
            if scanned:
                return {"ok": 0, "reason": "scanned_pdf", "fetched_at": now, "meta": meta}
            if len(smart_tokenize(text)) < 80:
                return {"ok": 0, "reason": "too_little_text", "fetched_at": now, "meta": meta}
            return {"ok": 1, "reason": None, "fetched_at": now, "meta": meta, "http_status": raw["status"],
                    "final_url": raw.get("final_url")}

        raw_text = ""
        if data:
            try:
                raw_text = raw_html_text(data.decode(raw.get("charset") or "utf-8", "replace"))
            except (LookupError, Exception):
                raw_text = raw_html_text(data.decode("utf-8", "replace"))
        raw_low = data[:200000].decode("utf-8", "ignore").lower()
        meta["raw_blocked"] = bool(raw.get("status") in (401, 403, 429, 503) and
                                   (any(m in raw_low for m in BLOCK_MARKERS) or CHALLENGE_RE.search(raw_text[:3000] or "")))
        try:
            rendered, shot, html = await self.render(url)
        except Exception as e:
            return {"ok": 0, "reason": f"render: {type(e).__name__}: {str(e)[:150]}", "fetched_at": now, "meta": meta}
        blocks = rendered.pop("blocks")
        text = "\n".join(b["text"] for b in blocks)
        meta.update(kind="html", render={k: v for k, v in rendered.items()}, final_url=rendered["final_url"],
                    http_status=rendered["status"])
        write_gz_json(cache_path(url, "json.gz"), {"kind": "html", "blocks": blocks, "raw_text": raw_text, "meta": meta})
        p = cache_path(url, "jpg")
        p.write_bytes(shot)
        with gzip.open(cache_path(url, "html.gz"), "wt", encoding="utf-8") as f:
            f.write(html)
        status = rendered["status"]
        if status and status >= 400:
            return {"ok": 0, "reason": f"http_{status}", "fetched_at": now, "meta": meta}
        head = (rendered.get("title") or "") + " " + text[:1500]
        if len(smart_tokenize(text)) < 300 and CHALLENGE_RE.search(head):
            return {"ok": 0, "reason": "challenge_in_browser", "fetched_at": now, "meta": meta}
        if len(smart_tokenize(text)) < 80:
            return {"ok": 0, "reason": "too_little_text", "fetched_at": now, "meta": meta}
        return {"ok": 1, "reason": None, "fetched_at": now, "meta": meta, "http_status": status,
                "final_url": rendered["final_url"]}


async def run(urls, concurrency):
    import json
    from playwright.async_api import async_playwright
    conn = db()
    done = [0]
    sem = asyncio.Semaphore(concurrency)
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(args=["--disable-blink-features=AutomationControlled"])
        limits = httpx.Limits(max_connections=concurrency * 4)
        async with httpx.AsyncClient(http2=True, follow_redirects=True, headers={"User-Agent": RAW_UA},
                                     verify=False, limits=limits, max_redirects=10) as client:
            fetcher = Fetcher(browser, client)

            async def one(url):
                async with sem:
                    try:
                        res = await asyncio.wait_for(fetcher.fetch(url), timeout=120)
                    except Exception as e:
                        res = {"ok": 0, "reason": f"timeout/{type(e).__name__}", "meta": {},
                               "fetched_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")}
                    conn.execute("INSERT OR REPLACE INTO fetches (url, ok, reason, final_url, http_status, fetched_at, meta) "
                                 "VALUES (?,?,?,?,?,?,?)", (url, res["ok"], res["reason"], res.get("final_url"),
                                                           res.get("http_status"), res["fetched_at"],
                                                           json.dumps(res["meta"], ensure_ascii=False)))
                    conn.commit()
                    done[0] += 1
                    if done[0] % 50 == 0 or done[0] == len(urls):
                        ok = conn.execute("SELECT sum(ok), count(*) FROM fetches").fetchone()
                        print(f"fetched {done[0]}/{len(urls)} (all-time ok {ok[0]}/{ok[1]})", file=sys.stderr, flush=True)
                    if not browser.is_connected():
                        raise RuntimeError("browser died")

            await asyncio.gather(*(one(u) for u in urls))
        await browser.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--concurrency", type=int, default=8)
    ap.add_argument("--urls-from", type=Path, help="file with one URL per line (default: all unfetched candidates)")
    ap.add_argument("--refetch", action="store_true", help="fetch even if already fetched")
    ap.add_argument("--source", help="only candidates from this source (e.g. search)")
    ap.add_argument("--added-since", help="only candidates added at or after this UTC timestamp")
    ap.add_argument("--max-ok", type=int, default=0, help="do nothing once this many pages (all-time) fetched ok")
    args = ap.parse_args()
    conn = db()
    ok = conn.execute("SELECT count(*) FROM fetches WHERE ok = 1").fetchone()[0]
    if args.max_ok and ok >= args.max_ok:
        print(f"{ok} pages fetched ok >= --max-ok {args.max_ok}; skipping", file=sys.stderr)
        return
    if args.urls_from:
        urls = [u.strip() for u in args.urls_from.read_text().splitlines() if u.strip()]
    else:
        urls = [r[0] for r in conn.execute("SELECT url FROM candidates WHERE (? IS NULL OR source = ?) AND (? IS NULL OR added_at >= ?)",
                                           [args.source, args.source, args.added_since, args.added_since])]
    if not args.refetch:
        done = {r[0] for r in conn.execute("SELECT url FROM fetches")}
        urls = [u for u in urls if u not in done]
    random.Random(7).shuffle(urls)
    if args.limit:
        urls = urls[:args.limit]
    print(f"fetching {len(urls)} urls", file=sys.stderr)
    import warnings
    warnings.filterwarnings("ignore")
    asyncio.run(run(urls, args.concurrency))


if __name__ == "__main__":
    main()
